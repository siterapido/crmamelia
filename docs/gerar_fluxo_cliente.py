#!/usr/bin/env python3
"""Gera fluxo-agente-cliente.docx — fluxograma para cliente final e time comercial."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return h


def add_flow_box(doc, title: str, lines: list[str], color: str = "E8F4FC"):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.size = Pt(11)
    for line in lines:
        p.add_run("  " + line + "\n").font.size = Pt(10)
    doc.add_paragraph()


def add_arrow(doc, label: str = ""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("▼ " + label if label else "▼")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x44, 0x72, 0xC4)


def add_decision(doc, question: str, branches: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "FFF4CE")
    p = cell.paragraphs[0]
    r = p.add_run("◆ " + question + "\n")
    r.bold = True
    for cond, out in branches:
        p.add_run(f"  • {cond} → {out}\n").font.size = Pt(10)
    doc.add_paragraph()


def add_table(doc, headers: list[str], rows: list[list[str]], header_color: str = "D6EAF8"):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
        set_cell_shading(t.rows[0].cells[j], header_color)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = val


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Como funciona a Amélia\n")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Assistente virtual de atendimento — Amélia Saúde\n").font.size = Pt(14)
    sub.add_run("Documento para clientes e equipe comercial").italic = True
    doc.add_page_break()

    add_heading(doc, "1. O que é a Amélia?", 1)
    doc.add_paragraph(
        "A Amélia é a assistente virtual da Amélia Saúde. Ela conversa com você pelo WhatsApp "
        "ou pelo chat do site, com tom de consultora acolhedora: tira dúvidas com base nas "
        "informações oficiais do site e organiza seus dados para um consultor humano montar "
        "a melhor proposta de plano de saúde."
    )

    add_heading(doc, "2. Fluxo da conversa", 1)
    add_flow_box(
        doc,
        "Passo 1 — Você entra em contato",
        [
            "WhatsApp ou chat no site (botão Quero Contratar)",
            "A Amélia responde em poucos instantes",
        ],
        "D5F5E3",
    )
    add_arrow(doc)
    add_flow_box(
        doc,
        "Passo 2 — Acolhimento",
        [
            "Ela entende seu motivo com empatia",
            "Não é um formulário rígido — é conversa",
        ],
    )
    add_arrow(doc)
    add_flow_box(
        doc,
        "Passo 3 — Informações e dúvidas",
        [
            "Pode explicar benefícios, operadoras (Nova Saúde, Ônix, Hapvida Notre Dame)",
            "Responde perguntas frequentes do site",
            "Referência de preço só como 'a partir de R$ 82' — valor final com consultor",
        ],
        "FDEBD0",
    )
    add_arrow(doc)
    add_decision(
        doc,
        "Ainda precisa de algum dado seu?",
        [
            ("Sim", "Uma pergunta por vez, na ordem natural da conversa"),
            ("Não", "Classifica prioridade e avisa o consultor"),
        ],
    )
    add_arrow(doc, "dados completos ou pedido de humano")
    add_flow_box(
        doc,
        "Passo 4 — Consultor humano",
        [
            "Especialista assume pelo WhatsApp ou CRM",
            "Monta proposta e valores para seu perfil",
        ],
        "D4E6F1",
    )

    add_heading(doc, "3. O que a Amélia pergunta (sem ordem fixa)", 2)
    for item in [
        "Seu nome",
        "Se o plano é para você/família (CPF) ou empresa (CNPJ)",
        "Cidade e estado",
        "Quantas pessoas no plano",
        "Se já tem plano e o que acha dele",
        "Para quando precisa do plano",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "4. O que a Amélia faz e não faz", 2)
    add_table(
        doc,
        ["Pode", "Não pode"],
        [
            ("Explicar tipos de plano e operadoras parceiras", "Fechar contrato sozinha"),
            ("Usar FAQ e benefícios do site", "Inventar preço ou promoção"),
            ("Organizar dados para o consultor", "Garantir rede hospitalar por cidade sem validação"),
            ("Encaminhar quando você pedir uma pessoa", "Substituir negociação de proposta"),
        ],
    )

    add_heading(doc, "5. Perguntas frequentes", 2)
    faqs = [
        (
            "Posso falar com uma pessoa?",
            "Sim. A qualquer momento peça na conversa e a Amélia encaminha para o time.",
        ),
        (
            "A Amélia define o valor do meu plano?",
            "Não. Ela pode citar a referência do site (a partir de R$ 82). O consultor confirma o valor final.",
        ),
        (
            "Onde vejo mais informações?",
            "Site ameliasaude.com.br, seção Como funciona a Amélia, FAQ e WhatsApp (21) 97233-8589.",
        ),
    ]
    for q, a in faqs:
        p = doc.add_paragraph()
        p.add_run(q + "\n").bold = True
        p.add_run(a)

    add_heading(doc, "6. Canais de contato", 2)
    doc.add_paragraph("WhatsApp: (21) 97233-8589 — segunda a sexta, 9h às 18h", style="List Bullet")
    doc.add_paragraph("0800: 0800-000-5123", style="List Bullet")
    doc.add_paragraph("Site: https://ameliasaude.com.br/lp", style="List Bullet")

    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "fluxo-agente-cliente.docx")
    doc.save(out_path)
    print(f"Gerado: {out_path}")
    print("PDF: abra no Word/LibreOffice e exporte como PDF, ou use: pandoc fluxo-agente-cliente.docx -o fluxo-agente-cliente.pdf")


if __name__ == "__main__":
    main()
