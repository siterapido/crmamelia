#!/usr/bin/env python3
"""Gera fluxo-agente-amelia.docx a partir do código em /Volumes/SSDdoMarcos/Projetos/crm-amelia."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return h


def add_flow_box(doc, title: str, lines: list[str], color: str = "E8F4FC"):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.size = Pt(11)
    for line in lines:
        r2 = p.add_run("  " + line + "\n")
        r2.font.size = Pt(10)
    doc.add_paragraph()


def add_arrow(doc, label: str = ""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("▼ " + label if label else "▼")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x44, 0x72, 0xC4)


def add_decision(doc, question: str, branches: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "FFF4CE")
    p = cell.paragraphs[0]
    r = p.add_run("◆ " + question + "\n")
    r.bold = True
    for cond, out in branches:
        p.add_run(f"  • {cond} → {out}\n").font.size = Pt(10)
    doc.add_paragraph()


def add_table(doc, headers: list[str], rows: list[list[str]], header_color: str = "D6EAF8"):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
        set_cell_shading(t.rows[0].cells[j], header_color)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = val


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Capa
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Fluxo do Agente Amélia\n")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("SDR WhatsApp — CRM Amélia\n").font.size = Pt(14)
    sub.add_run(
        "Documento gerado a partir do código em app/api/whatsapp e lib/ai\n"
    ).italic = True
    doc.add_page_break()

    add_heading(doc, "1. O que é este agente", 1)
    doc.add_paragraph(
        "O agente principal do WhatsApp é a Amélia SDR: um assistente de vendas que qualifica "
        "leads com metodologia SPIN Selling (6 perguntas), grava dados no CRM, move negócios no "
        "funil e, ao concluir, faz handoff para consultor humano. A IA usa Vercel AI SDK "
        "(generateObject) com OpenRouter e modelo moonshotai/kimi-k2.6."
    )
    doc.add_paragraph(
        "Outros fluxos de IA no mesmo repositório (site/chat/CMS) usam streamText e não fazem "
        "parte deste pipeline WhatsApp."
    )

    add_heading(doc, "2. Pontos de entrada", 2)
    add_table(
        doc,
        ["Entrada", "Arquivo / rota", "Papel"],
        [
            ("Webhook Evolution", "POST /api/whatsapp/webhook", "Fluxo principal inbound + status"),
            ("Diagnóstico", "GET /api/whatsapp/diagnostics", "Health check (Evolution, OpenRouter, DB)"),
            ("Teste E2E", "GET /api/whatsapp/test-agent", "Simula SDR; exige Bearer CRON_SECRET"),
            ("Follow-ups agendados", "GET /api/cron/followups", "Cron diário; envia mensagens da fila"),
            ("Inatividade", "POST /api/cron/inactivity", "Cron; recuperação e arquivo de conversas"),
            ("Chat site", "POST /api/chat/support", "streamText — não é o SDR WhatsApp"),
        ],
    )

    add_heading(doc, "3. Fluxograma — mensagem inbound (código atual)", 1)
    doc.add_paragraph(
        "Diferente de um webhook que só responde 200 e processa tudo depois: aqui a "
        "persistência no banco é síncrona antes do 200. A IA roda em after() somente após "
        "persistir com sucesso."
    )

    add_flow_box(
        doc,
        "1 — Cliente → Evolution → Webhook",
        [
            "WhatsApp envia mensagem à instância EVOLUTION_INSTANCE_NAME (ex.: amelia1)",
            "Evolution POST em /api/whatsapp/webhook",
            "Evento normalizado: MESSAGES_UPSERT → messages.upsert",
        ],
        "D5F5E3",
    )
    add_arrow(doc)

    add_flow_box(
        doc,
        "2 — Validação e resolução de contato",
        [
            "collectUpsertMessages: suporta data em array ou objeto único",
            "Ignora key.fromMe (mensagens enviadas pela instância)",
            "resolveContactFromUpsert: candidatos remoteJid, remoteJidAlt, sender, participant",
            "Exige JID respondível (@s.whatsapp.net ou @c.us)",
            "Extrai phone (dígitos) e replyJid para envio",
        ],
    )
    add_arrow(doc)

    add_decision(
        doc,
        "Contato resolvido?",
        [
            ("Não (LID @lid sem alt)", "skipped: lid_unresolved — sem CRM nem IA"),
            ("Não (outro)", "skipped: contact_unresolved"),
            ("Sim", "persistInboundMessage (síncrono)"),
        ],
    )
    add_arrow(doc, "sim")

    add_flow_box(
        doc,
        "3 — persistInboundMessage (ANTES do HTTP 200)",
        [
            "Upsert contacts (phone, whatsappId=replyJid, source=whatsapp)",
            "ensureDeal → deal no estágio pipeline new",
            "Busca/cria conversations status=active, aiEnabled default true",
            "INSERT messages inbound (sender: contact)",
            "Foto de perfil: updateProfilePictureAsync em background",
        ],
        "D4E6F1",
    )
    add_arrow(doc)

    add_decision(
        doc,
        "aiEnabled na conversa?",
        [
            ("false", "200 ok — só CRM; handoff humano ou operador"),
            ("true", "after() → runSDRAgent"),
        ],
    )
    add_arrow(doc, "sim + 200 JSON { status, persisted, skipped? }")

    add_flow_box(
        doc,
        "4 — runSDRAgent (assíncrono, maxDuration 60s)",
        [
            "Contato status=new → contacted + update_stage no funil",
            "sendPresence(phone, composing)",
            "Últimas 20 mensagens + buildSDRPrompt(contact, history)",
            "processSDRMessage → generateObject (Zod: reply + actions)",
            "executeSDRActions (qualify, stage, handoff, score, followup)",
            "sendAndSaveReply(replyJid) → normalizePhone → Evolution sendText",
            "INSERT message outbound sender=ai",
        ],
        "FDEBD0",
    )
    add_arrow(doc)

    add_flow_box(
        doc,
        "5 — Fim ou fallback",
        [
            "Sucesso: cliente recebe texto da Amélia no WhatsApp",
            "Erro IA/envio: FALLBACK_MESSAGE (instabilidade → humano)",
            "messages.update: keyId atualiza status sent/delivered/read no CRM",
        ],
        "D5F5E3",
    )

    add_heading(doc, "4. Motor de IA (lib/ai/sdr-agent.ts)", 2)
    add_table(
        doc,
        ["Parâmetro", "Valor"],
        [
            ("SDK", "Vercel AI SDK — generateObject (não streamText, sem ToolLoop)"),
            ("Provedor", "OpenRouter via @ai-sdk/openai"),
            ("Modelo", "moonshotai/kimi-k2.6"),
            ("Timeout", "20 segundos (Promise.race)"),
            ("Saída", "{ reply: string (min 1), actions: SDRAction[] }"),
            ("Auditoria", "INSERT ai_interactions (tokens, action, resumos)"),
        ],
    )

    add_heading(doc, "5. Fases de qualificação (lib/ai/sdr-prompt.ts)", 2)
    add_table(
        doc,
        ["Fase", "Condição", "Comportamento da Amélia"],
        [
            ("INICIO", "0 perguntas respondidas", "Apresentação + primeira pergunta (nome)"),
            ("COLETANDO_DADOS", "1–2 respondidas", "Uma pergunta por vez (SPIN)"),
            ("QUALIFICANDO", "3–5 respondidas", "Continua qualificação"),
            ("DADOS_COMPLETOS", "6/6 respondidas", "Agradecimento + handoff + score_lead"),
        ],
    )
    doc.add_paragraph(
        "As 6 perguntas: nome, perfil CPF/CNPJ, cidade/estado, vidas, plano atual, urgência. "
        "Cada resposta pode gerar action qualify; ao completar, score 1–5 e handoff."
    )

    add_heading(doc, "6. Ações do agente (lib/ai/sdr-actions.ts)", 2)
    add_table(
        doc,
        ["type", "Efeito no CRM"],
        [
            ("qualify", "Atualiza campo do contato/deal; auto-qualified se nome+endereço+vidas"),
            ("update_stage", "contacts.status + moveDealToStage (new, contacted, qualified…)"),
            ("handoff", "aiEnabled=false; status proposal; deal em proposta"),
            ("score_lead", "leadScore 1–5 + nota em deal/activities"),
            ("schedule_followup", "INSERT contact_followups → cron followups envia depois"),
        ],
    )

    add_heading(doc, "7. Agentes complementares (crons)", 2)
    add_flow_box(
        doc,
        "Cron followups — /api/cron/followups",
        [
            "Vercel Cron diário (vercel.json)",
            "Busca contact_followups com sent=false e scheduledAt <= agora",
            "sendTextMessage(contact.phone) via Evolution",
            "Grava mensagem outbound se houver conversationId",
        ],
        "E8DAEF",
    )
    add_flow_box(
        doc,
        "Cron inatividade — /api/cron/inactivity",
        [
            "POST com Bearer CRON_SECRET",
            "Monitora conversas active/dormant sem resposta do lead",
            "Envia RECOVERY_MESSAGES em thresholds configurados",
            "Pode arquivar conversa e markDealLostByInactivity",
        ],
        "E8DAEF",
    )

    add_heading(doc, "8. Ramificações e erros", 2)
    add_table(
        doc,
        ["Situação", "Comportamento"],
        [
            ("Todas persistências falham", "HTTP 500 persist_failed"),
            ("Algumas falham", "200 com persisted parcial + logs"),
            ("Instance mismatch", "Warn no log; processa mesmo assim"),
            ("Env EVOLUTION/OPENROUTER ausente", "Log CRITICAL; IA não roda ou falha no envio"),
            ("Evolution timeout", "5s por chamada HTTP"),
            ("Handoff executado", "Próximas mensagens: aiEnabled false — sem runSDRAgent"),
        ],
    )

    add_heading(doc, "9. Arquivos do repositório", 2)
    add_table(
        doc,
        ["Caminho", "Responsabilidade"],
        [
            ("app/api/whatsapp/webhook/route.ts", "Orquestração completa WhatsApp"),
            ("lib/ai/sdr-agent.ts", "generateObject + timeout + ai_interactions"),
            ("lib/ai/sdr-prompt.ts", "Prompt SPIN + fases + JSON de ações"),
            ("lib/ai/sdr-actions.ts", "CRM: qualify, stage, handoff, score, followup"),
            ("lib/ai/inactivity-messages.ts", "Textos do cron de recuperação"),
            ("lib/whatsapp/evolution-client.ts", "sendText, presence, profile picture"),
            ("lib/whatsapp/client.ts", "normalizePhone (JID/LID → dígitos)"),
            ("lib/db/schema.ts", "contacts, conversations, messages, deals, ai_interactions"),
            ("app/api/cron/followups/route.ts", "Disparo de follow-ups agendados"),
            ("app/api/cron/inactivity/route.ts", "Recuperação por inatividade"),
        ],
    )

    add_heading(doc, "10. Variáveis de ambiente obrigatórias", 2)
    for e in [
        "EVOLUTION_API_URL, EVOLUTION_API_KEY, EVOLUTION_INSTANCE_NAME",
        "OPENROUTER_API_KEY",
        "DATABASE_URL",
        "CRON_SECRET (test-agent, crons)",
    ]:
        doc.add_paragraph(e, style="List Bullet")

    add_heading(doc, "11. Diagrama resumido (texto)", 2)
    seq = """WhatsApp
  → Evolution API (webhook messages.upsert)
  → normalizeWebhookEvent + resolveContactFromUpsert
  → [sync] persistInboundMessage (contato, conversa, msg inbound)
  → HTTP 200 { persisted, skipped? }
  → [after, se aiEnabled] runSDRAgent
       → buildSDRPrompt + generateObject (OpenRouter)
       → executeSDRActions
       → sendTextMessage (Evolution)
       → msg outbound no CRM
  → [paralelo] messages.update → status entregue/lida"""
    p = doc.add_paragraph()
    run = p.add_run(seq)
    run.font.name = "Consolas"
    run.font.size = Pt(9)

    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "fluxo-agente-amelia.docx")
    doc.save(out_path)
    print(f"Gerado: {out_path}")


if __name__ == "__main__":
    main()
